import os
import shutil
import json
import csv
import zipfile
import traceback
from io import BytesIO
import pandas as pd
from fastapi import HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime
from sqlalchemy.orm import Session
from models.dashboard import Project
from schemas.dashboard import ProjectCreate, ProjectUpdate
from fastapi import HTTPException, UploadFile
from models.processing import ProcessingStatus, StatusEnum


# Define the base directory for the projects
PROJECTS_BASE_DIR = "projects"
# Define the colors for missing data (red) and present address (green)
RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
GREEN_FILL = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")


def create_project(db: Session, project: ProjectCreate):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    project_dir = os.path.join(PROJECTS_BASE_DIR, timestamp)

    try:
        create_project_directories(project_dir)

        db_project = Project(
            name=project.name,
            description=project.description,
            directory_name=timestamp,
        )
        db.add(db_project)
        db.commit()
        db.refresh(db_project)

        db_status = ProcessingStatus(
            project_id=db_project.id,
            status=StatusEnum.PENDING,
            total_files=0
        )
        db.add(db_status)
        db.commit()
        db.refresh(db_status)
        print(f"Project {db_project.name} committed to the database with status {db_status.status}.")

    except Exception as e:
        db.rollback()
        print(f"Failed to create project: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to create project or directories")

    return db_project

def create_project_directories(project_dir: str):
    previous_umask = os.umask(0)
    try:
        os.makedirs(project_dir, mode=0o777, exist_ok=True)
        os.makedirs(os.path.join(project_dir, "input"), mode=0o777, exist_ok=True)
        os.makedirs(os.path.join(project_dir, "output", "fail"), mode=0o777, exist_ok=True)
        os.makedirs(os.path.join(project_dir, "output", "success"), mode=0o777, exist_ok=True)
    finally:
        os.umask(previous_umask)

def get_projects(db: Session, skip: int = 0, limit: int = 10):
    return db.query(Project).offset(skip).limit(limit).all()

def update_project(db: Session, project_id: int, project: ProjectUpdate):
    db_project = db.query(Project).filter(Project.id == project_id).first()
    if db_project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    db_project.name = project.name
    db_project.description = project.description
    db.commit()
    db.refresh(db_project)
    return db_project

def delete_project(db: Session, project_id: int):
    db_project = db.query(Project).filter(Project.id == project_id).first()
    if db_project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    project_dir = os.path.join(PROJECTS_BASE_DIR, db_project.directory_name)
    try:
        db.delete(db_project)
        db.commit()
        print(f"Project {db_project.name} deleted from database.")
        delete_directory(project_dir)
    except Exception as e:
        print(f"Exception occurred during directory deletion: {str(e)}")
        db.rollback()
        raise HTTPException(status_code=500, detail="Failed to delete project or directory")
    return db_project

def delete_directory(directory_path: str):
    if os.path.exists(directory_path):
        try:
            shutil.rmtree(directory_path)
            print(f"Directory {directory_path} removed successfully.")
        except Exception as e:
            print(f"Failed to delete directory: {str(e)}")
            raise HTTPException(status_code=500, detail=f"Failed to delete directory: {str(e)}")
    else:
        print(f"Directory {directory_path} does not exist or was already removed.")

def upload_files_to_input(db: Session, project_id: int, files):
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    input_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "input")
    os.makedirs(input_dir, exist_ok=True)

    uploaded_files = save_files_to_directory(files, input_dir)
    return uploaded_files

def upload_folder_to_input(db: Session, project_id: int, files: list[UploadFile]):
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    input_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "input")
    os.makedirs(input_dir, mode=0o777, exist_ok=True)

    uploaded_files = save_files_to_directory(files, input_dir)
    return uploaded_files

def save_files_to_directory(files, directory: str):
    uploaded_files = []
    for file in files:
        file_path = os.path.join(directory, file.filename)
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        uploaded_files.append(file.filename)
    return uploaded_files

def delete_input_file(db: Session, project_id: int, filename: str):
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    input_file_path = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "input", filename)
    if not os.path.exists(input_file_path):
        raise HTTPException(status_code=404, detail="File not found")

    try:
        os.remove(input_file_path)
        return f"File '{filename}' successfully deleted from input directory."
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to delete file: {str(e)}")

def download_output_files(db: Session, project_id: int, convert_to_csv: bool = False):
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    output_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "output")
    if not os.path.exists(output_dir):
        raise HTTPException(status_code=404, detail="Output directory not found")

    zip_filepath = create_zip_with_conversion(output_dir, convert_to_csv)
    return zip_filepath

def download_success_output_files(db: Session, project_id: int, convert_to_csv: bool = False):
    """
    Download the successful output files for a given project.
    If convert_to_csv is True, it will convert all JSON files into a single CSV.
    """
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    output_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "output", "success")
    if not os.path.exists(output_dir):
        raise HTTPException(status_code=404, detail="Success output directory not found")

    if convert_to_csv:
        # Create a single CSV file from all JSON files
        csv_filepath = create_combined_csv(output_dir)
        return csv_filepath
    else:
        # Find all JSON files in the success directory
        json_files = [os.path.join(output_dir, f) for f in os.listdir(output_dir) if f.endswith('.json')]
        if not json_files:
            raise HTTPException(status_code=404, detail="No JSON files found")

        # If there is only one JSON file, return its path
        if len(json_files) == 1:
            return json_files[0]
        else:
            # Handle multiple JSON files by returning the first one (modify as needed)
            return json_files[0]

def create_combined_csv(db: Session, project_id: int, output_type: str) -> str:
    """
    Creates a combined CSV file from multiple JSON files in the specified output type directory.
    """
    # Fetch the project directory
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Define the directory and CSV path
    output_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "output", output_type)
    combined_csv_path = os.path.join(output_dir, "combined_success_output.csv")

    # Process JSON files and collect data
    all_data, all_fields = process_json_files(output_dir)

    # Create the combined CSV file
    try:
        with open(combined_csv_path, 'w', newline='', encoding='utf-8') as csv_file:
            writer = csv.DictWriter(csv_file, fieldnames=all_fields)
            writer.writeheader()
            for row in all_data:
                # Fill missing values with None (or an empty string if preferred)
                filled_row = {field: row.get(field, None) for field in all_fields}
                writer.writerow(filled_row)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to write combined CSV file: {str(e)}")

    return combined_csv_path


def process_json_files(output_dir: str) -> tuple[list[dict], list[str]]:
    """
    Processes all JSON files in the given directory and returns a list of flattened data
    and a list of all fields (columns) found across the JSON files.
    
    :param output_dir: The directory where JSON files are located
    :return: A tuple containing a list of dictionaries (flattened data) and a sorted list of all fields (column names)
    """
    json_files = [f for f in os.listdir(output_dir) if f.endswith('.json')]
    all_data = []
    all_fields = set()

    for json_file in json_files:
        json_file_path = os.path.join(output_dir, json_file)
        try:
            with open(json_file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                if isinstance(data, dict):
                    data = [data]  # Ensure data is always a list
                flattened_data = [flatten_json(item) for item in data]
                
                if flattened_data:
                    for row in flattened_data:
                        all_fields.update(row.keys())  # Collect all fields from all files
                    all_data.extend(flattened_data)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Failed to process JSON file {json_file}: {str(e)}")

    return all_data, sorted(all_fields)

def flatten_json(y):
    out = {}

    def flatten(x, name=''):
        if isinstance(x, dict):
            for a in x:
                flatten(x[a], name + a + '_')
        elif isinstance(x, list):
            i = 0
            for a in x:
                flatten(a, name + str(i) + '_')
                i += 1
        else:
            out[name[:-1]] = x

    flatten(y)
    return out


def download_fail_output_files(db: Session, project_id: int, convert_to_csv: bool = False):
    try:
        print(f"Attempting to download fail output files for project {project_id}")
        result = download_specific_output_files(db, project_id, "fail", convert_to_csv)
        print(f"Successfully downloaded fail output files for project {project_id}")
        return result
    except Exception as e:
        # Print a detailed traceback to the console
        print("Error occurred while downloading fail output files:")
        traceback.print_exc()  # This will print the full traceback of the exception
        raise HTTPException(status_code=500, detail="Internal Server Error")


def download_specific_output_files(db: Session, project_id: int, output_type: str, convert_to_csv: bool):
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    output_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "output", output_type)
    if not os.path.exists(output_dir):
        raise HTTPException(status_code=404, detail=f"{output_type.capitalize()} output directory not found")

    zip_filepath = create_zip_with_conversion(output_dir, convert_to_csv)
    return zip_filepath

def create_zip_with_conversion(directory: str, convert_to_csv: bool):
    zip_filename = os.path.basename(directory) + "_output.zip"
    zip_filepath = os.path.join(directory, zip_filename)

    # Allow both .csv and .json files to be zipped, regardless of conversion flag
    files_to_zip = [f for f in os.listdir(directory) if f.endswith('.csv') or f.endswith('.json')]
    print(f"Files selected for zipping: {files_to_zip}")

    if not files_to_zip:
        print(f"No files to zip in directory: {directory}")
        raise HTTPException(status_code=404, detail="No valid files to zip")

    with zipfile.ZipFile(zip_filepath, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as zipf:
        for file in files_to_zip:
            file_path = os.path.join(directory, file)
            if convert_to_csv and file.endswith(".json"):
                csv_filepath = file_path.replace(".json", ".csv")
                convert_json_to_csv(file_path, csv_filepath)
                file_to_zip = csv_filepath
            else:
                file_to_zip = file_path
            zipf.write(file_to_zip, os.path.relpath(file_to_zip, directory))
    return zip_filepath



def convert_json_to_csv(json_filepath, csv_filepath):
    try:
        with open(json_filepath, 'r') as json_file:
            data = json.load(json_file)
        
        if isinstance(data, dict):
            data = [data]  # Convert to list for consistent processing

        if not data:  # Check for empty JSON content
            raise ValueError("JSON content is empty or not in the expected format.")

        flattened_data = [flatten_json(item) for item in data]
        
        with open(csv_filepath, 'w', newline='') as csv_file:
            writer = csv.DictWriter(csv_file, fieldnames=flattened_data[0].keys())
            writer.writeheader()
            writer.writerows(flattened_data)
    except Exception as e:
        print(f"Failed to convert JSON to CSV: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to convert JSON to CSV: {str(e)}")

def get_file_tree(db: Session, project_id: int):
    """
    Retrieves the file tree for a specific project, sorted by file names.
    Only includes JSON files in the output sections.
    """
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    input_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "input")
    output_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "output")

    file_tree = {
        "input": get_directory_structure(input_dir),
        "output": {
            "success": get_directory_structure(os.path.join(output_dir, "success"), only_json=True),
            "fail": get_directory_structure(os.path.join(output_dir, "fail"), only_json=True),
        },
    }

    return file_tree


def get_directory_structure(rootdir: str, only_json: bool = False) -> dict:
    """
    Recursively constructs a dictionary representing the directory structure.
    The structure will be sorted by file and directory names.
    If only_json is True, only JSON files will be included in the structure.
    """
    structure = {}

    # Get a sorted list of all items in the directory
    items = sorted(os.listdir(rootdir), key=lambda s: s.lower())  # Sort case-insensitively

    for item in items:
        item_path = os.path.join(rootdir, item)
        if os.path.isdir(item_path):
            # Recursively get the directory structure for subdirectories
            structure[item] = get_directory_structure(item_path, only_json=only_json)
        else:
            # If only_json is True, skip non-JSON files
            if only_json and not item.endswith('.json'):
                continue
            # Mark the item as a file (or None, if you want a simpler output)
            structure[item] = None

    return structure

def delete_all_input_files(db: Session, project_id: int):
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    input_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "input")
    if not os.path.exists(input_dir):
        raise HTTPException(status_code=404, detail="Input directory not found")

    try:
        for filename in os.listdir(input_dir):
            file_path = os.path.join(input_dir, filename)
            if os.path.isfile(file_path):
                os.remove(file_path)
        return "All input files deleted successfully."
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to delete files: {str(e)}")

def delete_input_file_or_folder(db: Session, project_id: int, path: str):
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    input_path = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "input", path)
    if not os.path.exists(input_path):
        raise HTTPException(status_code=404, detail="File or folder not found")

    try:
        if os.path.isfile(input_path):
            os.remove(input_path)
        elif os.path.isdir(input_path):
            shutil.rmtree(input_path)
        return f"'{path}' successfully deleted from input directory."
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to delete '{path}': {str(e)}")

def generate_excel_for_project(db: Session, project_id: int, output_type: str = "success") -> BytesIO:
    """
    Generates an Excel file for a specific project, marking cells red where data is missing, 
    green where address information is complete, and yellow for "unbekannt". 
    Also adjusts column width based on the largest string in each column.
    """
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Define the output directory and ensure it exists
    output_dir = os.path.join(PROJECTS_BASE_DIR, project.directory_name, "output", output_type)
    os.makedirs(output_dir, exist_ok=True)

    # Process JSON files and collect data
    all_data, all_fields = process_json_files(output_dir)

    if not all_data:
        raise HTTPException(status_code=404, detail="No data available")

    try:
        # Create a DataFrame from the data
        df = pd.DataFrame(all_data, columns=all_fields)

        # Use a BytesIO stream to hold the Excel data
        output = BytesIO()
        writer = pd.ExcelWriter(output, engine='xlsxwriter')
        df.to_excel(writer, index=False, sheet_name='Project Data')

        workbook = writer.book
        worksheet = writer.sheets['Project Data']

        # Define formats
        red_format = workbook.add_format({'bg_color': '#FF0000'})
        green_format = workbook.add_format({'bg_color': '#00FF00'})
        yellow_format = workbook.add_format({'bg_color': '#FFFF00'})

        # Apply conditional formatting for missing data (red), address fields (green), and "unbekannt" (yellow)
        for idx, col in enumerate(df.columns):
            col_letter = get_column_letter(idx + 1)
            col_range = f'{col_letter}2:{col_letter}{len(df)+1}'

            # Conditional formatting for missing data
            worksheet.conditional_format(col_range, {
                'type': 'blanks',
                'format': red_format
            })
            worksheet.conditional_format(col_range, {
                'type': 'text',
                'criteria': 'equal to',
                'value': '""',
                'format': red_format
            })

            # Conditional formatting for "unbekannt"
            worksheet.conditional_format(col_range, {
                'type': 'text',
                'criteria': 'containing',
                'value': 'unbekannt',
                'format': yellow_format
            })

            # Conditional formatting for address-related fields
            if 'address' in col.lower():
                worksheet.conditional_format(col_range, {
                    'type': 'no_blanks',
                    'format': green_format
                })

            # Adjust column width based on the largest string in each column
            max_length = max(df[col].astype(str).map(len).max(), len(col)) + 2
            worksheet.set_column(f"{col_letter}:{col_letter}", max_length)

        writer.close()
        output.seek(0)  # Reset the pointer to the beginning of the stream

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Failed to generate Excel file: {str(e)}")

    return output

def upload_prompt_md(db: Session, project_id: int, md_file: UploadFile):
    """
    Save / replace a markdown prompt for the project.
    The file is always stored as <project_dir>/prompt.md
    """
    if not md_file.filename.endswith(".md"):
        raise HTTPException(status_code=400, detail="Only .md files allowed")

    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")

    project_root = os.path.join(PROJECTS_BASE_DIR, project.directory_name)
    os.makedirs(project_root, exist_ok=True)

    target_path = os.path.join(project_root, "prompt.md")

    # remove previous prompt (if any)
    if os.path.exists(target_path):
        os.remove(target_path)

    # copy the uploaded file
    with open(target_path, "wb") as buffer:
        shutil.copyfileobj(md_file.file, buffer)

    # record filename in DB (always "prompt.md")
    project.prompt_md = "prompt.md"
    db.commit()
    db.refresh(project)

    return {"message": "Prompt markdown uploaded", "filename": project.prompt_md}